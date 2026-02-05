import os
import json
import pandas as pd
from PyPDF2 import PdfReader
from pptx import Presentation
from docx import Document

# --------------------------------
# PATH CONFIGURATION
# --------------------------------
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(BASE_DIR, "Data")
OUTPUT_FILE = os.path.join(DATA_DIR, "processed", "all_documents.json")

documents = []

# --------------------------------
# LOADERS
# --------------------------------
def load_csv_or_excel(file_path):
    try:
        if file_path.lower().endswith(".csv"):
            df = pd.read_csv(file_path)

        elif file_path.lower().endswith(".xlsx"):
            df = pd.read_excel(file_path)

        else:
            # Unsupported file type
            return ""

        return df.astype(str).apply(" ".join, axis=1).str.cat(sep="\n")

    except Exception as e:
        print(f"Error reading {file_path}: {e}")
        return ""

def load_json(file_path):
    with open(file_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    # If it's a list of objects (like Nike products)
    if isinstance(data, list):
        blocks = []
        for item in data:
            if isinstance(item, dict):
                text = "\n".join(
                    f"{key}: {value}" for key, value in item.items()
                )
                blocks.append(text)
        return "\n\n".join(blocks)

    # Fallback
    return json.dumps(data, indent=2)


def load_pdf(file_path):
    reader = PdfReader(file_path)
    text = []
    for page in reader.pages:
        if page.extract_text():
            text.append(page.extract_text())
    return "\n".join(text)

def load_pptx(file_path):
    prs = Presentation(file_path)
    text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text.append(shape.text)
    return "\n".join(text)

def load_docx(file_path):
    doc = Document(file_path)
    return "\n".join([para.text for para in doc.paragraphs])

def load_emails(file_path):
    with open(file_path, "r", encoding="utf-8") as f:
        threads = json.load(f)

    blocks = []

    for thread_id, emails in threads.items():
        blocks.append(f"Thread ID: {thread_id}")

        for mail in emails:
            subject = mail.get("Subject", "")
            sender = mail.get("From", "")
            recipient = mail.get("To", "")
            date = mail.get("Date", "")
            body = mail.get("Body", "")

            blocks.append(
                f"From: {sender}\n"
                f"To: {recipient}\n"
                f"Date: {date}\n"
                f"Subject: {subject}\n"
                f"{body}"
            )

    return "\n\n".join(blocks)
# --------------------------------
# INGESTION FUNCTION
# --------------------------------
def ingest_folder(folder_name, loader, file_extensions, doc_type):
    folder_path = os.path.join(DATA_DIR, folder_name)

    if not os.path.exists(folder_path):
        return  # Safely skip missing folders

    for file in os.listdir(folder_path):
        if file.lower().endswith(file_extensions):
            file_path = os.path.join(folder_path, file)
            content = loader(file_path)

            # ✅ Only append if content is not empty
            if content.strip():  
                documents.append({
                    "source": file_path,
                    "type": doc_type,
                    "content": content
                })

# --------------------------------
# MAIN
# --------------------------------
def main():
    ingest_folder("excel", load_csv_or_excel, (".csv", ".xlsx"), "table")
    ingest_folder("json", load_json, (".json",), "json")
    ingest_folder("pdfs", load_pdf, (".pdf",), "pdf")
    ingest_folder("pptx", load_pptx, (".pptx",), "pptx")
    ingest_folder("doc", load_docx, (".docx",), "doc")

    # ✅ Process emails safely
    email_file = os.path.join(DATA_DIR, "emails", "threaded_emails.json")
    if os.path.exists(email_file):
        email_content = load_emails(email_file)
        if email_content.strip():  # Skip empty emails
            documents.append({
                "source": email_file,
                "type": "email",
                "content": email_content
            })

    # ✅ Save all documents
    os.makedirs(os.path.dirname(OUTPUT_FILE), exist_ok=True)
    with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
        json.dump(documents, f, indent=2)

    print(f"Data ingestion completed. {len(documents)} documents processed.")
    
if __name__ == "__main__":
    main()
